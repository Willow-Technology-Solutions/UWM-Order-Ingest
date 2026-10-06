"""
Helper functions for Dwelling Blocks → ERP PDF Automation workbook updates.

Row-append formatting follows the same openpyxl template-copy pattern used in
the Encompass Project helpers (style inheritance + table resize).
"""

from __future__ import annotations

from contextlib import contextmanager
from copy import copy
from datetime import date, datetime
from pathlib import Path
from shutil import copy2, move
from zoneinfo import ZoneInfo
import calendar
import ipaddress
import json
import os
import re
import socket
import sys
import time
from typing import Any

import requests
from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from logger import setupLogger


def getBaseDirectory() -> Path:
    """Return the directory that contains the bundled executable or script."""
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


baseDirPath = getBaseDirectory()
exportsDirPath = baseDirPath / "exports"
archiveDirPath = baseDirPath / "archive"
bundledEnvPath = Path(__file__).resolve().parent / ".env"
scriptEnvPath = baseDirPath / ".env"


def loadEnvironmentFile() -> tuple[Path, Exception | None]:
    """Load the runtime `.env`, falling back to the bundled copy when needed.

    Returns the path that was loaded, and any error from copying a bundled file.
    Exits when no `.env` file exists.
    """
    if scriptEnvPath.exists():
        load_dotenv(scriptEnvPath)
        return scriptEnvPath, None

    if bundledEnvPath.exists():
        load_dotenv(bundledEnvPath)
        try:
            copy2(bundledEnvPath, scriptEnvPath)
        except Exception as error:
            return bundledEnvPath, error
        return bundledEnvPath, None

    sys.stderr.write(
        f"The .env file was not found at {scriptEnvPath} or bundled location "
        f"{bundledEnvPath}. Please create it and try again.\n"
    )
    sys.exit(1)


_loadedEnvPath, _bundledCopyError = loadEnvironmentFile()
logger = setupLogger()
logger.info(f"Base directory resolved to: {baseDirPath}")
logger.info(f"Loaded environment variables from {_loadedEnvPath}")
if _bundledCopyError is not None:
    logger.error(
        f"Failed to copy bundled .env to {scriptEnvPath}: {_bundledCopyError}"
    )

ERP_PDF_AUTOMATION_XLSX_PATH = os.getenv("ERP_PDF_AUTOMATION_XLSX_PATH")
GSG_CONNECT_BASE_URL = os.getenv("GSG_CONNECT_BASE_URL")
GSG_CONNECT_USERNAME = os.getenv("GSG_CONNECT_USERNAME")
GSG_CONNECT_PASSWORD = os.getenv("GSG_CONNECT_PASSWORD")
COMPANY_ID = (os.getenv("COMPANY_ID") or "").strip()
FILE_CLEANUP_DAYS_THRESHOLD = int(os.getenv("FILE_CLEANUP_DAYS_THRESHOLD") or "7")
HTTP_TIMEOUT = float(os.getenv("HTTP_TIMEOUT") or "60")
DNS_SERVER = (os.getenv("DNS_SERVER") or "").strip() or None
CURRENT_ORDERS_ORDER_DATE_MONTH_LOOKBACK = int(
    os.getenv("CURRENT_ORDERS_ORDER_DATE_MONTH_LOOKBACK") or "6"
)
ERP_SHEET_NAME = "ERP"


def _isIpAddressLiteral(host: str | bytes) -> bool:
    if isinstance(host, bytes):
        try:
            host = host.decode("ascii")
        except UnicodeDecodeError:
            return False
    try:
        ipaddress.ip_address(host)
        return True
    except ValueError:
        return False


@contextmanager
def _socket_getaddrinfo_replaced(replacement):
    previous = socket.getaddrinfo
    socket.getaddrinfo = replacement
    try:
        yield
    finally:
        socket.getaddrinfo = previous


def _doh_resolve_a_records(hostname: str, *, dns_server: str, system_getaddrinfo) -> list[str]:
    """Resolve A records via Cloudflare-style JSON DoH (used by 1.1.1.1)."""
    with _socket_getaddrinfo_replaced(system_getaddrinfo):
        try:
            response = requests.get(
                f"https://{dns_server}/dns-query",
                params={"name": hostname, "type": "A"},
                headers={
                    "Accept": "application/dns-json",
                    "Host": "cloudflare-dns.com",
                },
                timeout=HTTP_TIMEOUT,
            )
            response.raise_for_status()
            payload = response.json()
        except Exception as error:
            logger.warning(
                f"DNS override lookup failed for {hostname} using {dns_server}: {error}"
            )
            return []
    answers = payload.get("Answer", [])
    return [
        item["data"]
        for item in answers
        if item.get("type") == 1 and item.get("data")
    ]


def _pinned_getaddrinfo_factory(*, dns_server: str, system_getaddrinfo):
    def pinned(host, port, family=0, type=0, proto=0, flags=0):
        if not host:
            return system_getaddrinfo(host, port, family, type, proto, flags)
        name = host.decode() if isinstance(host, bytes) else str(host)
        if _isIpAddressLiteral(name):
            return system_getaddrinfo(host, port, family, type, proto, flags)
        ips = _doh_resolve_a_records(
            name, dns_server=dns_server, system_getaddrinfo=system_getaddrinfo
        )
        if ips:
            sock_type = type or socket.SOCK_STREAM
            proto_val = proto or socket.IPPROTO_TCP
            return [(socket.AF_INET, sock_type, proto_val, "", (ip, port)) for ip in ips]
        return system_getaddrinfo(host, port, family, type, proto, flags)

    return pinned


class DnsPinnedSession(requests.Session):
    """Session that resolves names through DNS-over-HTTPS when dns_server is set."""

    def __init__(self, *, dns_server: str):
        super().__init__()
        self.trust_env = False
        self._system_getaddrinfo = socket.getaddrinfo
        self._pinned = _pinned_getaddrinfo_factory(
            dns_server=dns_server, system_getaddrinfo=self._system_getaddrinfo
        )

    def request(self, method, url, **kwargs):
        with _socket_getaddrinfo_replaced(self._pinned):
            return super().request(method, url, **kwargs)


def _createDnsPinnedSession(*, dns_server: str | None = None) -> requests.Session:
    if dns_server:
        return DnsPinnedSession(dns_server=dns_server)
    return requests.Session()

# Columns that must stay blank in new Dwelling Blocks rows.
INTENTIONALLY_BLANK_COLUMNS = {
    "File Number",
    "Client Reference Number",
    "Additonal Information",
    "Report Uploaded",
    "Appointment Date",
}

AUTOMATION_INTEGER_COLUMNS = {"ID", "Loan Number", "Street #", "Client Reference Number"}
AUTOMATION_NUMERIC_COLUMNS = {"Total Payout Amount"}
AUTOMATION_DATE_COLUMNS = {
    "Order Date",
    "Due Date",
    "Report Uploaded",
    "Appointment Date",
}
AUTOMATION_PHONE_COLUMNS = {
    "Borrower's Phone Number",
    "Co-Borrower's Phone Number",
    "Access Contact's Phone Number",
}
AUTOMATION_HYPERLINK_COLUMNS = {
    "Borrower's Email Address",
    "Co-Borrower's Email Address",
}

DEFAULT_DATE_NUMBER_FORMAT = "m/d/yy h:mm"
DEFAULT_INTEGER_NUMBER_FORMAT = "0"
DEFAULT_CURRENCY_NUMBER_FORMAT = '"$"#,##0.00'
DEFAULT_ZIP_NUMBER_FORMAT = "@"
LOCAL_TIMEZONE = ZoneInfo("America/New_York")


def resolveAutomationWorkbookPath() -> str:
    """
    Resolve the ERP PDF Automation workbook path from environment or the base directory.
    """
    logger.info(f"ERP_PDF_AUTOMATION_XLSX_PATH: {ERP_PDF_AUTOMATION_XLSX_PATH}")
    if ERP_PDF_AUTOMATION_XLSX_PATH and Path(ERP_PDF_AUTOMATION_XLSX_PATH).exists():
        return ERP_PDF_AUTOMATION_XLSX_PATH

    output_path = baseDirPath / "ERP PDF Automation.xlsx"
    if output_path.exists():
        return str(output_path)

    raise FileNotFoundError(
        "Automation workbook not found. Set ERP_PDF_AUTOMATION_XLSX_PATH to a valid "
        "file path, or place 'ERP PDF Automation.xlsx' next to the script."
    )


def backupAutomationWorkbook(workbook_path) -> Path | None:
    """Create a timestamped backup copy of the automation workbook before processing."""
    workbook_path = Path(workbook_path)
    if not workbook_path.exists():
        logger.warning(f"Automation workbook not found for backup: {workbook_path}")
        return None

    archiveDirPath.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup_path = (
        archiveDirPath / f"{workbook_path.stem} Backup {timestamp}{workbook_path.suffix}"
    )
    copy2(workbook_path, backup_path)
    logger.info(f"Created automation workbook backup at {backup_path}")
    return backup_path


def sendSuccessHealthcheck() -> None:
    """Send a success ping to Healthcheck.io when configured."""
    healthcheck_url = os.getenv("HEALTHCHECK_IO_URL")
    if not healthcheck_url:
        logger.warning("HEALTHCHECK_IO_URL not set, skipping health check ping")
        return

    try:
        session = _createDnsPinnedSession(dns_server=DNS_SERVER)
        response = session.get(healthcheck_url, timeout=HTTP_TIMEOUT)
        if response.status_code == 200:
            logger.info("Successfully sent ping to Healthcheck.io")
        else:
            logger.warning(
                f"Healthcheck.io ping returned status code {response.status_code}"
            )
    except Exception as e:
        logger.exception(f"Error sending ping to Healthcheck.io: {e}")


def archiveFile(file_path) -> None:
    """Archive a file after successful processing."""
    if not file_path or not os.path.exists(file_path):
        logger.warning(f"File path is empty or does not exist: {file_path}")
        return

    try:
        source_path = Path(file_path)
        archiveDirPath.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        archived_path = (
            archiveDirPath / f"{source_path.stem} {timestamp}{source_path.suffix}"
        )
        move(str(source_path), str(archived_path))
        logger.info(f"Successfully archived {file_path} to {archived_path}")
    except Exception as e:
        logger.error(f"Failed to archive {file_path}: {e}")


def cleanupOldFiles(days_threshold: int = FILE_CLEANUP_DAYS_THRESHOLD) -> None:
    """Clean up old files from the exports and archive directories."""
    cutoff_time = time.time() - (days_threshold * 86400)

    for directory in [exportsDirPath, archiveDirPath]:
        if not directory.exists():
            continue
        for file_path in directory.glob("*"):
            if file_path.is_file() and file_path.stat().st_mtime < cutoff_time:
                try:
                    file_path.unlink()
                    logger.info(f"Deleted old file: {file_path}")
                except Exception as e:
                    logger.error(f"Failed to delete old file {file_path}: {e}")


def exportOrders(
    *,
    view: str = "in-progress",
    page_size: int = 100,
    max_results: int | None = None,
) -> list[dict[str, str]]:
    """Pull Dwelling Blocks assignments and map them to ERP column rows."""
    from dwelling_blocks import (
        VIEW_FILTERS,
        DwellingBlocksClient,
        assignmentToErpRow,
        enrichAssignmentWithDetail,
        loadCredentials,
        loadErpColumns,
    )

    columns = loadErpColumns()
    client_id, client_secret = loadCredentials()
    client = DwellingBlocksClient(client_id, client_secret)

    logger.info("Authenticating to Dwelling Blocks...")
    client.getAccessToken()
    logger.info("Access token acquired.")

    filters = VIEW_FILTERS[view]
    logger.info(
        "Pulling assignments via POST /vendors/orders/search "
        "(view=%s, page_size=%s)...",
        view,
        page_size,
    )
    orders = list(
        client.iterVendorOrders(
            page_size=page_size,
            max_results=max_results,
            use_search=True,
            filters=filters,
        )
    )
    orders = filterOrdersByLookback(orders)
    logger.info(
        "Enriching %s assignment(s) with contacts/properties/transfers...",
        len(orders),
    )
    enriched = [enrichAssignmentWithDetail(client, item) for item in orders]
    return [assignmentToErpRow(order, columns) for order in enriched]


def processOrdersFile(
    rows: list[dict[str, Any]], output_file: str | Path
) -> tuple[int, list[Any]]:
    """
    Append exported ERP rows to the automation workbook, matching existing cell formatting.

    Skips loans already present in GSG Connect for COMPANY_ID.
    Each new row ID is the loan number with COMPANY_ID appended, so the same
    order keeps the same ID if it is written again.

    Returns:
        tuple[int, list]: Number of new orders added and their loan numbers
    """
    output_file = Path(output_file)
    if not output_file.exists():
        logger.error(f"Automation workbook not found: {output_file}")
        return 0, []

    if not rows:
        logger.info("No ERP rows to append")
        return 0, []

    if not COMPANY_ID:
        logger.error(
            "COMPANY_ID is not set; unable to dedupe against past orders. "
            "Skipping order processing."
        )
        return 0, []
    if _companyIdDigits(COMPANY_ID) is None:
        logger.error(
            "COMPANY_ID has no digits; unable to build stable row IDs. "
            "Skipping order processing."
        )
        return 0, []

    authenticated_client = _authenticateToAsp(
        base_url=GSG_CONNECT_BASE_URL,
        username=GSG_CONNECT_USERNAME,
        password=GSG_CONNECT_PASSWORD,
    )
    if not authenticated_client:
        logger.error(
            "Unable to authenticate to GSG Connect API; skipping order processing"
        )
        return 0, []

    past_orders_payload = fetchPastOrdersPayload(authenticated_client)
    if past_orders_payload is None:
        logger.error(
            "Unable to load past orders from GSG Connect API; skipping order processing"
        )
        return 0, []

    existing_loan_company_pairs = _loanCompanyPairsFromPastOrdersPayload(
        past_orders_payload
    )
    seen_loan_company_pairs = set(existing_loan_company_pairs)
    logger.info(
        "Deduping Dwelling Blocks rows against Company_ID=%s (%s existing pairs loaded)",
        COMPANY_ID,
        len(seen_loan_company_pairs),
    )

    try:
        workbook = load_workbook(output_file)
        worksheet = _getErpWorksheet(workbook)
        headers = _readWorksheetHeaders(worksheet)
        if not headers:
            logger.error(f"No headers found in {output_file}")
            return 0, []
        if "Loan Number" not in headers:
            logger.error(f"'Loan Number' column not found in {output_file}")
            return 0, []

        template_row_index = _findAutomationTemplateRow(worksheet, headers)
        column_template_row_lookup = _findAutomationColumnTemplateRows(
            worksheet,
            headers,
            fallback_row_index=template_row_index,
        )

        added_loan_numbers: list[Any] = []
        added_count = 0

        for row in rows:
            loan_number = _normalizeLoanNumber(row.get("Loan Number"))
            if loan_number is None:
                logger.info("Skipping row with missing Loan Number")
                continue
            if (loan_number, COMPANY_ID) in seen_loan_company_pairs:
                continue

            outgoing = dict(row)
            for blank_column in INTENTIONALLY_BLANK_COLUMNS:
                outgoing[blank_column] = None

            if not outgoing.get("ID"):
                outgoing["ID"] = _buildAutomationRowId(loan_number, COMPANY_ID)

            if not outgoing.get("File Name"):
                outgoing["File Name"] = _buildFileName(outgoing)

            if not outgoing.get("FHA Case Number"):
                outgoing["FHA Case Number"] = "Not set"

            country = outgoing.get("Country")
            if country and str(country).strip().upper() in {"US", "USA", "UNITED STATES"}:
                outgoing["Country"] = "USA"

            _writeAutomationRow(
                worksheet,
                headers,
                outgoing,
                template_row_index=template_row_index,
                template_row_lookup=column_template_row_lookup,
            )
            seen_loan_company_pairs.add((loan_number, COMPANY_ID))
            added_loan_numbers.append(loan_number)
            added_count += 1

        if added_count:
            _extendWorksheetTablesToMaxRow(worksheet, headers)
            workbook.save(output_file)
            logger.info(f"Added {added_count} new orders to {output_file}")
            logger.info(
                "Loan numbers added to workbook: %s",
                ", ".join(str(loan_number) for loan_number in added_loan_numbers),
            )
            return added_count, added_loan_numbers

        logger.info("No new orders to add")
        return 0, []
    except Exception as e:
        logger.exception(f"Error appending ERP rows: {e}")
        return 0, []


def _getErpWorksheet(workbook):
    if ERP_SHEET_NAME in workbook.sheetnames:
        return workbook[ERP_SHEET_NAME]
    worksheet = workbook.active
    if worksheet is None:
        raise RuntimeError("Workbook has no active worksheet")
    return worksheet


def _readWorksheetHeaders(worksheet: Worksheet) -> list[str]:
    headers = [cell.value if cell.value is not None else "" for cell in worksheet[1]]
    while headers and headers[-1] == "":
        headers.pop()
    return headers


def _authenticateToAsp(
    *,
    base_url: str | None,
    username: str | None,
    password: str | None,
    timeout: float = HTTP_TIMEOUT,
) -> requests.Session | None:
    """
    Log into the legacy ASP site and return an authenticated session.

    Same approach as Encompass Project helpers._authenticateToAsp.
    """
    if not base_url:
        logger.error("GSG_CONNECT_BASE_URL is required for _authenticateToAsp.")
        return None
    if not username or not password:
        logger.error(
            "GSG_CONNECT_USERNAME and GSG_CONNECT_PASSWORD are required for "
            "_authenticateToAsp."
        )
        return None

    resolved_login_url = f"{base_url}/processlogin.asp"
    client = _createDnsPinnedSession(dns_server=DNS_SERVER)
    login_data = {
        "username": username,
        "userpassword": password,
    }

    try:
        logger.info(f"Posting ASP login request to {resolved_login_url}")
        response = client.post(
            resolved_login_url,
            data=login_data,
            timeout=timeout,
            allow_redirects=True,
        )
        if not (response.status_code == 200 and "default.asp" in response.url.lower()):
            logger.error(
                "Authentication failed. "
                f"Status code: {response.status_code}. Final URL: {response.url}"
            )
            return None
        logger.info(f"Authentication successful. Final URL: {response.url}")
        return client
    except requests.RequestException as error:
        logger.exception(f"Error logging into ASP site: {error}")
        return None


def fetchPastOrdersPayload(
    authenticated_client: requests.Session, timeout: float = HTTP_TIMEOUT
) -> list | None:
    """
    Fetch raw JSON rows from get_orders.asp (orders table snapshot).

    Same approach as Encompass Project helpers.fetchPastOrdersPayload.
    """
    if not GSG_CONNECT_BASE_URL:
        logger.error("GSG_CONNECT_BASE_URL is required for past orders lookup.")
        return None

    resolved_api_url = f"{GSG_CONNECT_BASE_URL}/json/get_orders.asp"
    try:
        logger.info(f"Fetching past orders from {resolved_api_url}")
        api_response = authenticated_client.get(resolved_api_url, timeout=timeout)
        api_response.raise_for_status()
        payload = json.loads(api_response.text, strict=False)

        if isinstance(payload, dict) and payload.get("error"):
            logger.error(f"Past orders API returned error: {payload!r}")
            return None
        if not isinstance(payload, list):
            logger.error(
                f"Past orders API returned unexpected JSON type: {type(payload)}"
            )
            return None
        return payload
    except requests.RequestException as exc:
        logger.exception(f"Error fetching past orders: {exc}")
        return None
    except json.JSONDecodeError as exc:
        logger.exception(f"Error parsing past orders JSON: {exc}")
        return None


def _loanCompanyPairsFromPastOrdersPayload(
    payload: list,
) -> set[tuple[int, str]]:
    """Build (Loan_Number, Company_ID) pairs from get_orders.asp rows."""
    pairs: set[tuple[int, str]] = set()
    for row in payload:
        if not isinstance(row, dict):
            continue
        loan_normalized = _normalizeLoanNumber(row.get("Loan_Number"))
        company_raw = row.get("Company_ID")
        if loan_normalized is None or company_raw is None:
            continue
        company_id = str(company_raw).strip()
        if not company_id:
            continue
        pairs.add((loan_normalized, company_id))

    logger.info(f"Loaded {len(pairs)} past order loan/client pairs from API")
    return pairs


def _companyIdDigits(company_id: str) -> str | None:
    """Return the digit characters of a company id, matching Encompass client ids."""
    digits = re.sub(r"\D", "", company_id.strip())
    return digits or None


def _buildAutomationRowId(loan_number: int, company_id: str) -> int:
    """Stable workbook ID: loan number with the company id appended.

    Same formula as Encompass ``{loan_number}{client_id}``.
    """
    company_digits = _companyIdDigits(company_id)
    if not company_digits:
        raise ValueError("COMPANY_ID has no digits")
    return int(f"{loan_number}{company_digits}")


def _buildFileName(row: dict[str, Any]) -> str:
    street_number = str(row.get("Street #") or "").strip()
    street_address = str(row.get("Street Address") or "").strip()
    if street_number and street_address:
        return f"{street_number} {street_address}.pdf"
    if street_address:
        return f"{street_address}.pdf"
    return ""


def _normalizeLoanNumber(value: Any) -> int | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value if value >= 0 else None
    if isinstance(value, float):
        if value.is_integer() and value >= 0:
            return int(value)
        return None

    text = str(value).strip().replace(",", "")
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    if number.is_integer() and number >= 0:
        return int(number)
    return None


def _normalizeNumericValue(value: Any) -> int | float | None | Any:
    if value is None:
        return None
    cleaned_value = value
    if isinstance(value, str):
        cleaned_value = value.strip().replace("$", "").replace(",", "")
        if cleaned_value.startswith("(") and cleaned_value.endswith(")"):
            cleaned_value = f"-{cleaned_value[1:-1]}"
        if not cleaned_value:
            return None
    try:
        numeric_value = float(str(cleaned_value))
    except (TypeError, ValueError):
        return value
    if numeric_value.is_integer():
        return int(numeric_value)
    return numeric_value


def _subtractMonths(value: datetime, months: int) -> datetime:
    """Move a datetime back by calendar months, clamping the day when needed."""
    month_index = value.year * 12 + (value.month - 1) - months
    year, month_zero = divmod(month_index, 12)
    month = month_zero + 1
    day = min(value.day, calendar.monthrange(year, month)[1])
    return value.replace(year=year, month=month, day=day)


def _lookbackCutoff(months: int) -> datetime:
    """Midnight Eastern on the oldest Order Date still inside the lookback."""
    today_eastern = datetime.now(LOCAL_TIMEZONE).replace(
        hour=0, minute=0, second=0, microsecond=0, tzinfo=None
    )
    return _subtractMonths(today_eastern, months)


def orderDateWithinLookback(order_date: Any, months: int | None = None) -> bool:
    """Return True when Order Date falls on or after the lookback cutoff."""
    lookback_months = (
        CURRENT_ORDERS_ORDER_DATE_MONTH_LOOKBACK if months is None else months
    )
    parsed = _normalizeDateValue(order_date)
    if not isinstance(parsed, datetime):
        return False
    order_day = parsed.replace(hour=0, minute=0, second=0, microsecond=0)
    return order_day >= _lookbackCutoff(lookback_months)


def filterOrdersByLookback(orders: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Keep assignments whose createdOn Order Date is inside the lookback window."""
    months = CURRENT_ORDERS_ORDER_DATE_MONTH_LOOKBACK
    cutoff = _lookbackCutoff(months)
    kept = [
        order
        for order in orders
        if orderDateWithinLookback(order.get("createdOn"), months)
    ]
    logger.info(
        "Current Orders: kept %s of %s rows with Order Date on or after %s "
        "(%s-month lookback)",
        len(kept),
        len(orders),
        cutoff.date(),
        months,
    )
    return kept


def _to_local_naive(value: datetime) -> datetime:
    """Convert aware datetimes to America/New_York naive (matches ERP sheet)."""
    if value.tzinfo is None:
        return value
    return value.astimezone(LOCAL_TIMEZONE).replace(tzinfo=None)


def _normalizeDateValue(value: Any) -> datetime | None | Any:
    """Parse date-like values, preserving time (Excel format m/d/yy h:mm)."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return _to_local_naive(value)
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())

    text = str(value).strip()
    if not text or text.upper() in {"N/A", "NOT SET", "NONE"}:
        return None

    # ISO-8601 with optional fractional seconds / Z
    iso_text = text.replace("Z", "+00:00")
    try:
        parsed = datetime.fromisoformat(iso_text)
        return _to_local_naive(parsed)
    except ValueError:
        pass

    for fmt in (
        "%m/%d/%Y %H:%M",
        "%m/%d/%y %H:%M",
        "%Y-%m-%d %H:%M:%S",
        "%m/%d/%Y",
        "%m/%d/%y",
    ):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return value


def _formatPhoneNumber(value: Any) -> str | None:
    if value is None:
        return None
    digits = re.sub(r"\D", "", str(value))
    if len(digits) == 11 and digits.startswith("1"):
        digits = digits[1:]
    if len(digits) == 10:
        return f"({digits[0:3]}) {digits[3:6]}-{digits[6:10]}"
    text = str(value).strip()
    return text or None


def _normalizeStreetNumber(value: Any) -> int | str | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    if re.fullmatch(r"\d+", text):
        return int(text)
    return text


def _coerceAutomationCellValue(header: str, value: Any) -> Any:
    """Coerce outgoing values to match existing Excel cell types."""
    if header in INTENTIONALLY_BLANK_COLUMNS:
        return None

    if value is None:
        return None

    if isinstance(value, str):
        value = value.strip()
        if not value:
            return None

    if header in AUTOMATION_PHONE_COLUMNS:
        return _formatPhoneNumber(value)

    if header == "Street #":
        return _normalizeStreetNumber(value)

    if header == "Zip":
        text = str(value).strip()
        return text or None

    if header in AUTOMATION_INTEGER_COLUMNS:
        if header == "ID":
            loan_like = _normalizeLoanNumber(value)
            return loan_like if loan_like is not None else value
        return _normalizeLoanNumber(value)

    if header in AUTOMATION_NUMERIC_COLUMNS:
        return _normalizeNumericValue(value)

    if header in AUTOMATION_DATE_COLUMNS:
        return _normalizeDateValue(value)

    return value


def _applyHyperlinkStyle(cell) -> None:
    """Apply the workbook's Hyperlink named style so email cells match existing rows."""
    workbook = cell.parent.parent
    if "Hyperlink" not in workbook.named_styles:
        return
    cell.style = "Hyperlink"


def _findAutomationTemplateRow(worksheet: Worksheet, headers: list[str]) -> int | None:
    """Return the first populated data row to use as a formatting template."""
    if worksheet.max_row < 2:
        return None

    for row_index in range(2, worksheet.max_row + 1):
        if any(
            worksheet.cell(row=row_index, column=column_index).value not in (None, "")
            for column_index in range(1, len(headers) + 1)
        ):
            return row_index
    return None


def _findAutomationColumnTemplateRows(
    worksheet: Worksheet,
    headers: list[str],
    fallback_row_index: int | None = None,
) -> dict[str, int | None]:
    """Return a per-column template row so each appended cell inherits the right format."""
    template_rows: dict[str, int | None] = {}
    if worksheet.max_row < 2:
        return {header: fallback_row_index for header in headers}

    for column_index, header in enumerate(headers, start=1):
        template_row_index = fallback_row_index
        for row_index in range(2, worksheet.max_row + 1):
            cell_value = worksheet.cell(row=row_index, column=column_index).value
            if cell_value not in (None, ""):
                template_row_index = row_index
                break
        template_rows[header] = template_row_index
    return template_rows


def _extendWorksheetTablesToMaxRow(worksheet: Worksheet, headers: list[str]) -> None:
    """Expand existing worksheet tables so appended rows stay inside the Excel table."""
    if not worksheet.tables:
        return

    table_ref = f"A1:{get_column_letter(len(headers))}{worksheet.max_row}"
    for table_name in list(worksheet.tables.keys()):
        worksheet.tables[table_name].ref = table_ref


def _writeAutomationRow(
    worksheet: Worksheet,
    headers: list[str],
    row_values: dict[str, Any],
    template_row_index: int | None = None,
    template_row_lookup: dict[str, int | None] | None = None,
) -> None:
    """Write a new automation row while preserving workbook formatting."""
    target_row_index = worksheet.max_row + 1

    for column_index, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=target_row_index, column=column_index)
        cell.value = _coerceAutomationCellValue(header, row_values.get(header))

        column_template_row_index = template_row_index
        if template_row_lookup is not None:
            column_template_row_index = template_row_lookup.get(header, template_row_index)

        if column_template_row_index is not None:
            template_cell = worksheet.cell(
                row=column_template_row_index, column=column_index
            )
            if template_cell.has_style:
                cell._style = copy(template_cell._style)

        if header in AUTOMATION_INTEGER_COLUMNS and isinstance(cell.value, int):
            cell.number_format = DEFAULT_INTEGER_NUMBER_FORMAT

        if header in AUTOMATION_NUMERIC_COLUMNS and isinstance(cell.value, (int, float)):
            if not cell.number_format or cell.number_format == "General":
                cell.number_format = DEFAULT_CURRENCY_NUMBER_FORMAT

        if header in AUTOMATION_DATE_COLUMNS and isinstance(cell.value, datetime):
            if not cell.number_format or cell.number_format == "General":
                cell.number_format = DEFAULT_DATE_NUMBER_FORMAT

        if header == "Zip" and cell.value is not None:
            cell.number_format = DEFAULT_ZIP_NUMBER_FORMAT

        if header in AUTOMATION_HYPERLINK_COLUMNS and cell.value not in (None, ""):
            _applyHyperlinkStyle(cell)
